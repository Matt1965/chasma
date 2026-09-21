#!/usr/bin/env python3
"""Author Slice 8 UAL animation + weapon/equipment presentation data."""

from __future__ import annotations

import openpyxl
from pathlib import Path

WORKBOOK = Path(__file__).resolve().parents[1] / "Chasma Design.xlsx"


def header_map(ws) -> dict[str, int]:
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    return {str(h).strip(): idx for idx, h in enumerate(headers) if h}


def ensure_column(ws, mapping: dict[str, int], name: str) -> int:
    if name in mapping:
        return mapping[name]
    col = ws.max_column
    ws.cell(row=1, column=col + 1, value=name)
    mapping[name] = col
    return col


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)

    profiles = wb["Animation Profiles"]
    pmap = header_map(profiles)
    for row in profiles.iter_rows(min_row=2):
        if row[pmap["Profile ID"]].value != "human_base":
            continue
        row[pmap["Idle Animation"]].value = "Idle"
        row[pmap["Walk Animation"]].value = "Walk"
        row[pmap["Run Animation"]].value = "Run"
        if "Work Animation" in pmap:
            row[pmap["Work Animation"]].value = "Mine"
        if "Death Animation" in pmap:
            row[pmap["Death Animation"]].value = "Death"
        if "Hit Reaction Animation" in pmap:
            row[pmap["Hit Reaction Animation"]].value = "Hit"
        if "Upper Body Split Bone" in pmap:
            row[pmap["Upper Body Split Bone"]].value = "spine_02"
        break

    weapons = wb["Weapons"]
    wmap = header_map(weapons)
    fam_col = ensure_column(weapons, wmap, "Animation Family")
    idle_col = ensure_column(weapons, wmap, "Combat Idle Clip")
    variant_col = ensure_column(weapons, wmap, "Attack Variant")
    anim_col = wmap["Animation Key"]

    weapon_updates = {
        "weapon_fists": ("Punch_Jab", "Unarmed", "", "Punch_Cross"),
        "weapon_iron_sword": ("Sword_Attack", "OneHandSword", "Sword_Idle", ""),
    }
    for row in weapons.iter_rows(min_row=2):
        wid = row[wmap["Weapon ID"]].value
        if wid not in weapon_updates:
            continue
        anim, family, combat_idle, variant = weapon_updates[wid]
        row[anim_col].value = anim
        row[fam_col].value = family
        row[idle_col].value = combat_idle
        row[variant_col].value = variant

    visuals = wb["Equipment Visuals"]
    vmap = header_map(visuals)
    for col in (
        "Stowed Socket",
        "Stowed Local Translation",
        "Stowed Local Rotation",
        "Stowed Local Scale",
    ):
        ensure_column(visuals, vmap, col)

    stowed = {
        "iron_sword": {
            "Stowed Socket": "RightHip",
            "Stowed Local Translation": "0.12,0.02,-0.08",
            "Stowed Local Rotation": "0.0,0.707,0.0,0.707",
        }
    }
    for row in visuals.iter_rows(min_row=2):
        item = row[vmap["Item ID"]].value
        if item not in stowed:
            continue
        for key, value in stowed[item].items():
            row[vmap[key]].value = value

    wb.save(WORKBOOK)
    print(f"Updated {WORKBOOK} for Slice 8")


if __name__ == "__main__":
    main()
