#!/usr/bin/env python3
"""Populate Equipment Visuals fit columns for skinned armor rows."""

from __future__ import annotations

from pathlib import Path

import openpyxl

WORKBOOK = Path(__file__).resolve().parents[1] / "Chasma Design.xlsx"
SHEET = "Equipment Visuals"

SKINNED_ITEMS = {
    "peasant_body",
    "peasant_arms",
    "peasant_legs",
    "ranger_body",
    "ranger_arms",
    "ranger_legs",
    "ranger_hood",
}

FIT_BY_ITEM = {
    "peasant_body": ("1.04", ""),
    "peasant_arms": ("1.035", ""),
    "peasant_legs": ("1.04", ""),
    "ranger_body": ("1.04", ""),
    "ranger_arms": ("1.035", ""),
    "ranger_legs": ("1.04", ""),
    "ranger_hood": ("1.03", "0,0.004,0"),
}


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb[SHEET]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    for name in ("Fit Scale", "Fit Offset"):
        if name not in headers:
            headers[name] = ws.max_column + 1
            ws.cell(1, headers[name], name)
    item_col = headers["Item ID"]
    mode_col = headers["Presentation Mode"]
    for row in range(2, ws.max_row + 1):
        item_id = (ws.cell(row, item_col).value or "").strip()
        mode = (ws.cell(row, mode_col).value or "").strip().lower()
        if item_id not in SKINNED_ITEMS or "skinned" not in mode:
            continue
        fit_scale, fit_offset = FIT_BY_ITEM[item_id]
        ws.cell(row, headers["Fit Scale"], fit_scale)
        ws.cell(row, headers["Fit Offset"], fit_offset)
    wb.save(WORKBOOK)
    print(f"updated fit columns on {SHEET}")


if __name__ == "__main__":
    main()
