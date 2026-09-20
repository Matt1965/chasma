#!/usr/bin/env python3
"""Update Equipment Visuals consumed morph params (CG5 + CG9 regional)."""

from __future__ import annotations

from pathlib import Path

import openpyxl

WORKBOOK = Path(__file__).resolve().parents[1] / "Chasma Design.xlsx"
SHEET = "Equipment Visuals"
COLUMN = "Consumed Morph Params"

BODY_PARAMS = "build,fat,muscle,shoulders,torso,hips"
ARMS_PARAMS = "build,fat,muscle,arms"
LEGS_PARAMS = "build,fat,muscle,legs,hips"
HEAD_PARAMS = "head_size"
EMPTY = ""

BODY_ITEMS = {"peasant_body", "ranger_body"}
ARMS_ITEMS = {"peasant_arms", "ranger_arms"}
LEGS_ITEMS = {"peasant_legs", "ranger_legs"}
HEAD_ITEMS = {"ranger_hood"}
FEET_ITEMS = {"peasant_feet", "ranger_feet"}


def consumed_for_item(item_id: str) -> str:
    if item_id in BODY_ITEMS:
        return BODY_PARAMS
    if item_id in ARMS_ITEMS:
        return ARMS_PARAMS
    if item_id in LEGS_ITEMS:
        return LEGS_PARAMS
    if item_id in HEAD_ITEMS:
        return HEAD_PARAMS
    if item_id in FEET_ITEMS:
        return EMPTY
    return EMPTY


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb[SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    header_map = {str(h).strip(): idx for idx, h in enumerate(headers) if h}
    if COLUMN not in header_map:
        header_map[COLUMN] = len(headers)
        ws.cell(row=1, column=header_map[COLUMN] + 1, value=COLUMN)

    item_col = header_map["Item ID"]
    consumed_col = header_map[COLUMN]
    for row in ws.iter_rows(min_row=2):
        item_id = str(row[item_col].value or "").strip()
        if not item_id:
            continue
        row[consumed_col].value = consumed_for_item(item_id)

    wb.save(WORKBOOK)
    print(f"updated {WORKBOOK} ({SHEET}.{COLUMN})")


if __name__ == "__main__":
    main()
