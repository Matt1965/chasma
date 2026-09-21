#!/usr/bin/env python3
"""Author Slice 7.1 real equipment content into Chasma Design.xlsx."""

from __future__ import annotations

import openpyxl
from pathlib import Path

WORKBOOK = Path(__file__).resolve().parents[1] / "Chasma Design.xlsx"

REMOVE_ITEM_IDS = {
    "simple_helmet",
    "padded_vest",
    "basic_arm_guards",
    "basic_leg_guards",
    "basic_boots",
    "scrap_sword",
    "basic_backpack",
}

REMOVE_ARMOR_PROFILES = {
    "armor_starter_head",
    "armor_starter_body",
    "armor_starter_arms",
    "armor_starter_legs",
    "armor_starter_feet",
}

REMOVE_WEAPON_IDS = {"weapon_scrap_sword"}

ITEMS = [
    # Peasant armor
    ("peasant_body", "Peasant Body", "armor", 3, 3, 1200, "body", None, "armor_peasant_body", None),
    ("peasant_arms", "Peasant Arms", "armor", 2, 2, 600, "arms", None, "armor_peasant_arms", None),
    ("peasant_legs", "Peasant Legs", "armor", 2, 3, 900, "legs", None, "armor_peasant_legs", None),
    ("peasant_feet", "Peasant Feet", "armor", 2, 2, 700, "feet", None, "armor_peasant_feet", None),
    # Ranger armor
    ("ranger_hood", "Ranger Hood", "armor", 2, 2, 800, "head", None, "armor_ranger_hood", None),
    ("ranger_body", "Ranger Body", "armor", 3, 3, 1200, "body", None, "armor_ranger_body", None),
    ("ranger_arms", "Ranger Arms", "armor", 2, 2, 600, "arms", None, "armor_ranger_arms", None),
    ("ranger_legs", "Ranger Legs", "armor", 2, 3, 900, "legs", None, "armor_ranger_legs", None),
    ("ranger_feet", "Ranger Feet", "armor", 2, 2, 700, "feet", None, "armor_ranger_feet", None),
    # Weapons
    ("iron_dagger", "Iron Dagger", "weapon", 1, 2, 500, "weapon", "weapon_iron_dagger", None, None),
    ("iron_hand_axe", "Iron Hand Axe", "weapon", 1, 2, 900, "weapon", "weapon_iron_hand_axe", None, None),
    ("iron_sword", "Iron Sword", "weapon", 1, 3, 1500, "weapon", "weapon_iron_sword", None, None),
    ("iron_greatsword", "Iron Greatsword", "weapon", 1, 4, 2200, "weapon", "weapon_iron_greatsword", None, None),
    ("iron_greataxe", "Iron Greataxe", "weapon", 2, 3, 2400, "weapon", "weapon_iron_greataxe", None, None),
    ("iron_warhammer", "Iron Warhammer", "weapon", 1, 4, 2100, "weapon", "weapon_iron_warhammer", None, None),
    ("wooden_bow", "Wooden Bow", "weapon", 1, 3, 900, "weapon", "weapon_wooden_bow", None, None),
    # Backpack
    ("leather_backpack", "Leather Backpack", "container", 2, 3, 500, "backpack", None, None, "backpack_basic_internal"),
]

ARMOR_PROFILES = [
    ("armor_peasant_body", "Peasant Body Armor", "Light peasant torso protection.", 5),
    ("armor_peasant_arms", "Peasant Arm Armor", "Light peasant arm protection.", 3),
    ("armor_peasant_legs", "Peasant Leg Armor", "Light peasant leg protection.", 4),
    ("armor_peasant_feet", "Peasant Foot Armor", "Light peasant foot protection.", 3),
    ("armor_ranger_hood", "Ranger Hood Armor", "Ranger head protection.", 5),
    ("armor_ranger_body", "Ranger Body Armor", "Ranger torso protection.", 15),
    ("armor_ranger_arms", "Ranger Arm Armor", "Ranger arm protection.", 5),
    ("armor_ranger_legs", "Ranger Leg Armor", "Ranger leg protection.", 10),
    ("armor_ranger_feet", "Ranger Foot Armor", "Ranger foot protection.", 5),
]

WEAPONS = [
    ("weapon_iron_dagger", "Iron Dagger", "Short iron blade.", 8, "Piercing", 1.5, 1.8, 0.12, 0.1, "Melee", None, "attack_fists", "Enemies"),
    ("weapon_iron_hand_axe", "Iron Hand Axe", "One-handed iron axe.", 11, "Slashing", 2.0, 1.2, 0.18, 0.12, "Melee", None, "attack_fists", "Enemies"),
    ("weapon_iron_sword", "Iron Sword", "Standard iron sword.", 14, "Slashing", 2.2, 1.0, 0.2, 0.15, "Melee", None, "attack_fists", "Enemies"),
    ("weapon_iron_greatsword", "Iron Greatsword", "Large two-handed iron sword.", 20, "Slashing", 2.8, 0.7, 0.28, 0.2, "Melee", None, "attack_fists", "Enemies"),
    ("weapon_iron_greataxe", "Iron Greataxe", "Heavy two-handed greataxe.", 22, "Slashing", 2.6, 0.6, 0.3, 0.22, "Melee", None, "attack_fists", "Enemies"),
    ("weapon_iron_warhammer", "Iron Warhammer", "Heavy iron warhammer.", 18, "Blunt", 2.4, 0.75, 0.26, 0.18, "Melee", None, "attack_fists", "Enemies"),
    ("weapon_wooden_bow", "Wooden Bow", "Simple wooden bow.", 12, "Piercing", 15.0, 0.8, 0.35, 0.25, "Projectile", "iron_arrow", "attack_fists", "Enemies"),
]

RIGID_ITEMS = [
    "iron_dagger",
    "iron_hand_axe",
    "iron_sword",
    "iron_greatsword",
    "iron_greataxe",
    "iron_warhammer",
    "wooden_bow",
    "leather_backpack",
]

SKINNED_ITEMS = [
    "peasant_body",
    "peasant_arms",
    "peasant_legs",
    "peasant_feet",
    "ranger_hood",
    "ranger_body",
    "ranger_arms",
    "ranger_legs",
    "ranger_feet",
]

TORSO_MORPH_PARAMS = "build,fat,muscle"
HEAD_MORPH_PARAMS = "head_size"

SKINNED_ASSET_NAMES = {
    "peasant_body": "peasant_body",
    "peasant_arms": "peasant_arms",
    "peasant_legs": "peasant_legs",
    "peasant_feet": "peasant_feet",
    "ranger_hood": "ranger_hood",
    "ranger_body": "ranger_body",
    "ranger_arms": "ranger_arms",
    "ranger_legs": "ranger_legs",
    "ranger_feet": "ranger_feet",
}


def header_map(ws) -> dict[str, int]:
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    return {str(h).strip(): idx for idx, h in enumerate(headers) if h}


def ensure_headers(ws, headers: list[str]) -> dict[str, int]:
    mapping = header_map(ws)
    for header in headers:
        if header not in mapping:
            mapping[header] = len(mapping)
            ws.cell(row=1, column=mapping[header] + 1, value=header)
    return mapping


def delete_rows_by_key(ws, mapping: dict[str, int], key_col: str, keys: set[str]) -> None:
    key_idx = mapping[key_col]
    rows_to_delete = []
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row=row, column=key_idx + 1).value
        if value in keys:
            rows_to_delete.append(row)
    for row in reversed(rows_to_delete):
        ws.delete_rows(row, 1)


def upsert_row(ws, mapping: dict[str, int], key_col: str, key: str, row_data: dict) -> None:
    key_idx = mapping[key_col]
    target = None
    for row in range(2, ws.max_row + 2):
        existing = ws.cell(row=row, column=key_idx + 1).value
        if existing == key:
            target = row
            break
        if existing in (None, ""):
            target = row
            break
    if target is None:
        target = ws.max_row + 1
    for header, value in row_data.items():
        if header not in mapping:
            continue
        ws.cell(row=target, column=mapping[header] + 1, value=value)


def build_equipment_visual_rows() -> list[dict]:
    rows: list[dict] = []
    for item_id in RIGID_ITEMS:
        render_key = item_id
        socket = "Back" if item_id == "leather_backpack" else "RightHand"
        for unit_key in ("human_male", "human_female"):
            rows.append(
                {
                    "Item ID": item_id,
                    "Unit Render Key": unit_key,
                    "Equipped Render Key": render_key,
                    "Presentation Mode": "RigidAttachment",
                    "Socket": socket,
                }
            )
    for item_id in SKINNED_ITEMS:
        asset = SKINNED_ASSET_NAMES[item_id]
        if item_id in {"ranger_hood"}:
            consumed = HEAD_MORPH_PARAMS
        elif item_id in {"ranger_feet", "peasant_feet"}:
            consumed = ""
        else:
            consumed = TORSO_MORPH_PARAMS
        for unit_key in ("human_male", "human_female"):
            rows.append(
                {
                    "Item ID": item_id,
                    "Unit Render Key": unit_key,
                    "Equipped Render Key": f"equipment/{unit_key}/{asset}",
                    "Presentation Mode": "SkinnedOverlay",
                    "Consumed Morph Params": consumed,
                }
            )
    return rows


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)

    items = wb["Items"]
    item_map = ensure_headers(
        items,
        [
            "Item ID",
            "Name",
            "Category",
            "Width",
            "Height",
            "Stackable",
            "Max Stack",
            "Mass Grams",
            "Enabled",
            "Description",
            "Base Value",
            "Unique Instance Required",
            "Equipment Slots",
            "Weapon Definition ID",
            "Armor Profile ID",
            "Backpack Profile ID",
            "Render Key",
        ],
    )
    delete_rows_by_key(items, item_map, "Item ID", REMOVE_ITEM_IDS)
    for (
        item_id,
        name,
        category,
        width,
        height,
        mass,
        slot,
        weapon_id,
        armor_id,
        backpack_id,
    ) in ITEMS:
        row = {
            "Item ID": item_id,
            "Name": name,
            "Category": category,
            "Width": width,
            "Height": height,
            "Stackable": "N",
            "Max Stack": 1,
            "Mass Grams": mass,
            "Enabled": "Y",
            "Description": name,
            "Base Value": 5,
            "Unique Instance Required": "Y",
            "Equipment Slots": slot,
        }
        if weapon_id:
            row["Weapon Definition ID"] = weapon_id
        if armor_id:
            row["Armor Profile ID"] = armor_id
        if backpack_id:
            row["Backpack Profile ID"] = backpack_id
        if item_id in RIGID_ITEMS:
            row["Render Key"] = item_id
        upsert_row(items, item_map, "Item ID", item_id, row)

    if "Armor Profiles" not in wb.sheetnames:
        armor_ws = wb.create_sheet("Armor Profiles")
        armor_ws.append(["Profile ID", "Name", "Description", "Armor Rating", "Enabled"])
    armor_ws = wb["Armor Profiles"]
    armor_map = ensure_headers(
        armor_ws,
        ["Profile ID", "Name", "Description", "Armor Rating", "Enabled"],
    )
    delete_rows_by_key(armor_ws, armor_map, "Profile ID", REMOVE_ARMOR_PROFILES)
    for profile_id, name, description, rating in ARMOR_PROFILES:
        upsert_row(
            armor_ws,
            armor_map,
            "Profile ID",
            profile_id,
            {
                "Profile ID": profile_id,
                "Name": name,
                "Description": description,
                "Armor Rating": rating,
                "Enabled": "Y",
            },
        )

    weapons = wb["Weapons"]
    weapon_map = ensure_headers(
        weapons,
        [
            "Weapon ID",
            "Name",
            "Description",
            "Damage",
            "Damage Type",
            "Range",
            "Attacks Per Second",
            "Windup",
            "Recovery",
            "Hit Mode",
            "Projectile Key",
            "Animation Key",
            "Target Filters",
            "Stat Scaling",
            "Enabled",
        ],
    )
    delete_rows_by_key(weapons, weapon_map, "Weapon ID", REMOVE_WEAPON_IDS)
    for (
        weapon_id,
        name,
        description,
        damage,
        damage_type,
        range_m,
        aps,
        windup,
        recovery,
        hit_mode,
        projectile_key,
        animation_key,
        target_filters,
    ) in WEAPONS:
        upsert_row(
            weapons,
            weapon_map,
            "Weapon ID",
            weapon_id,
            {
                "Weapon ID": weapon_id,
                "Name": name,
                "Description": description,
                "Damage": damage,
                "Damage Type": damage_type,
                "Range": range_m,
                "Attacks Per Second": aps,
                "Windup": windup,
                "Recovery": recovery,
                "Hit Mode": hit_mode,
                "Projectile Key": projectile_key,
                "Animation Key": animation_key,
                "Target Filters": target_filters,
                "Enabled": "Y",
            },
        )

    if "Equipment Visuals" in wb.sheetnames:
        del wb["Equipment Visuals"]
    visuals_ws = wb.create_sheet("Equipment Visuals")
    visual_headers = [
        "Item ID",
        "Unit Render Key",
        "Equipped Render Key",
        "Presentation Mode",
        "Socket",
        "Local Translation",
        "Local Rotation",
        "Local Scale",
        "Consumed Morph Params",
    ]
    visuals_ws.append(visual_headers)
    for row in build_equipment_visual_rows():
        visuals_ws.append(
            [
                row["Item ID"],
                row["Unit Render Key"],
                row["Equipped Render Key"],
                row["Presentation Mode"],
                row.get("Socket"),
                None,
                None,
                None,
                row.get("Consumed Morph Params"),
            ]
        )

    wb.save(WORKBOOK)
    print(f"Updated {WORKBOOK}")


if __name__ == "__main__":
    main()
