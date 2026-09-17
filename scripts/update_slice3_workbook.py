#!/usr/bin/env python3
"""Author Slice 3 starter equipment, armor profiles, and scrap sword into Chasma Design.xlsx."""

from __future__ import annotations

import openpyxl
from pathlib import Path

WORKBOOK = Path(__file__).resolve().parents[1] / "Chasma Design.xlsx"

EQUIPMENT_ITEMS = [
    {
        "Item ID": "simple_helmet",
        "Name": "Simple Helmet",
        "Category": "armor",
        "Width": 2,
        "Height": 2,
        "Stackable": "N",
        "Max Stack": 1,
        "Mass Grams": 800,
        "Enabled": "Y",
        "Description": "Basic head protection.",
        "Base Value": 5,
        "Unique Instance Required": "Y",
        "Equipment Slots": "head",
        "Armor Profile ID": "armor_starter_head",
    },
    {
        "Item ID": "padded_vest",
        "Name": "Padded Vest",
        "Category": "armor",
        "Width": 3,
        "Height": 3,
        "Stackable": "N",
        "Max Stack": 1,
        "Mass Grams": 1200,
        "Enabled": "Y",
        "Description": "Light torso padding.",
        "Base Value": 5,
        "Unique Instance Required": "Y",
        "Equipment Slots": "body",
        "Armor Profile ID": "armor_starter_body",
    },
    {
        "Item ID": "basic_arm_guards",
        "Name": "Basic Arm Guards",
        "Category": "armor",
        "Width": 2,
        "Height": 2,
        "Stackable": "N",
        "Max Stack": 1,
        "Mass Grams": 600,
        "Enabled": "Y",
        "Description": "Simple arm protection.",
        "Base Value": 5,
        "Unique Instance Required": "Y",
        "Equipment Slots": "arms",
        "Armor Profile ID": "armor_starter_arms",
    },
    {
        "Item ID": "basic_leg_guards",
        "Name": "Basic Leg Guards",
        "Category": "armor",
        "Width": 2,
        "Height": 3,
        "Stackable": "N",
        "Max Stack": 1,
        "Mass Grams": 900,
        "Enabled": "Y",
        "Description": "Simple leg protection.",
        "Base Value": 5,
        "Unique Instance Required": "Y",
        "Equipment Slots": "legs",
        "Armor Profile ID": "armor_starter_legs",
    },
    {
        "Item ID": "basic_boots",
        "Name": "Basic Boots",
        "Category": "armor",
        "Width": 2,
        "Height": 2,
        "Stackable": "N",
        "Max Stack": 1,
        "Mass Grams": 700,
        "Enabled": "Y",
        "Description": "Sturdy boots.",
        "Base Value": 5,
        "Unique Instance Required": "Y",
        "Equipment Slots": "feet",
        "Armor Profile ID": "armor_starter_feet",
    },
    {
        "Item ID": "scrap_sword",
        "Name": "Scrap Sword",
        "Category": "weapon",
        "Width": 1,
        "Height": 3,
        "Stackable": "N",
        "Max Stack": 1,
        "Mass Grams": 1500,
        "Enabled": "Y",
        "Description": "Crude blade for starter combat testing.",
        "Base Value": 5,
        "Unique Instance Required": "Y",
        "Equipment Slots": "weapon",
        "Weapon Definition ID": "weapon_scrap_sword",
    },
    {
        "Item ID": "basic_backpack",
        "Name": "Basic Backpack",
        "Category": "container",
        "Width": 2,
        "Height": 3,
        "Stackable": "N",
        "Max Stack": 1,
        "Mass Grams": 500,
        "Enabled": "Y",
        "Description": "Standard carried pack with internal storage.",
        "Base Value": 5,
        "Unique Instance Required": "Y",
        "Equipment Slots": "backpack",
        "Backpack Profile ID": "backpack_basic_internal",
    },
]

ARMOR_PROFILES = [
    ("armor_starter_head", "Starter Head Armor", "Provisional starter head armor profile.", 5),
    ("armor_starter_body", "Starter Body Armor", "Provisional starter body armor profile.", 15),
    ("armor_starter_arms", "Starter Arm Armor", "Provisional starter arm armor profile.", 5),
    ("armor_starter_legs", "Starter Leg Armor", "Provisional starter leg armor profile.", 10),
    ("armor_starter_feet", "Starter Foot Armor", "Provisional starter foot armor profile.", 5),
]

SCRAP_SWORD_WEAPON = [
    "weapon_scrap_sword",
    "Scrap Sword",
    "Crude blade for starter combat testing.",
    14.0,
    "Slashing",
    2.2,
    1.0,
    0.2,
    0.15,
    "Melee",
    None,
    "attack_fists",
    "Enemies",
    None,
    "Y",
]


def header_map(ws) -> dict[str, int]:
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    return {str(h).strip(): idx for idx, h in enumerate(headers) if h}


def ensure_headers(ws, headers: list[str]) -> dict[str, int]:
    mapping = header_map(ws)
    for header in headers:
        if header not in mapping:
            mapping[header] = len(mapping)
            ws.cell(row=1, column=mapping[header] + 1, value=header)
    return mapping


def upsert_row(ws, mapping: dict[str, int], key_col: str, key: str, row_data: dict) -> None:
    key_idx = mapping[key_col]
    for row in range(2, ws.max_row + 2):
        existing = ws.cell(row=row, column=key_idx + 1).value
        if existing == key:
            target = row
            break
        if existing in (None, ""):
            target = row
            break
    else:
        target = ws.max_row + 1
    for header, value in row_data.items():
        if header not in mapping:
            continue
        ws.cell(row=target, column=mapping[header] + 1, value=value)


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)

    categories = wb["Item Categories"]
    cat_map = ensure_headers(categories, ["Category ID", "Name", "Enabled", "Description", "Sort Priority"])
    upsert_row(
        categories,
        cat_map,
        "Category ID",
        "container",
        {
            "Category ID": "container",
            "Name": "Container",
            "Enabled": "Y",
            "Description": "Portable containers and packs",
            "Sort Priority": 120,
        },
    )

    items = wb["Items"]
    item_headers = [
        "Item ID",
        "Name",
        "Category",
        "Width",
        "Height",
        "Stackable",
        "Max Stack",
        "Mass Grams",
        "Enabled",
        "Description",
        "Base Value",
        "Unique Instance Required",
        "Equipment Slots",
        "Weapon Definition ID",
        "Armor Profile ID",
        "Backpack Profile ID",
    ]
    item_map = ensure_headers(items, item_headers)
    for row in EQUIPMENT_ITEMS:
        upsert_row(items, item_map, "Item ID", row["Item ID"], row)

    if "Armor Profiles" not in wb.sheetnames:
        armor_ws = wb.create_sheet("Armor Profiles")
        armor_ws.append(
            [
                "Profile ID",
                "Name",
                "Description",
                "Armor Rating",
                "Enabled",
            ]
        )
    armor_ws = wb["Armor Profiles"]
    armor_map = ensure_headers(
        armor_ws,
        [
            "Profile ID",
            "Name",
            "Description",
            "Armor Rating",
            "Enabled",
        ],
    )
    for profile_id, name, description, rating in ARMOR_PROFILES:
        upsert_row(
            armor_ws,
            armor_map,
            "Profile ID",
            profile_id,
            {
                "Profile ID": profile_id,
                "Name": name,
                "Description": description,
                "Armor Rating": rating,
                "Enabled": "Y",
            },
        )

    weapons = wb["Weapons"]
    weapon_map = header_map(weapons)
    weapon_row = {
        "Weapon ID": SCRAP_SWORD_WEAPON[0],
        "Name": SCRAP_SWORD_WEAPON[1],
        "Description": SCRAP_SWORD_WEAPON[2],
        "Damage": SCRAP_SWORD_WEAPON[3],
        "Damage Type": SCRAP_SWORD_WEAPON[4],
        "Range": SCRAP_SWORD_WEAPON[5],
        "Attacks Per Second": SCRAP_SWORD_WEAPON[6],
        "Windup": SCRAP_SWORD_WEAPON[7],
        "Recovery": SCRAP_SWORD_WEAPON[8],
        "Hit Mode": SCRAP_SWORD_WEAPON[9],
        "Projectile Key": SCRAP_SWORD_WEAPON[10],
        "Animation Key": SCRAP_SWORD_WEAPON[11],
        "Target Filters": SCRAP_SWORD_WEAPON[12],
        "Stat Scaling": SCRAP_SWORD_WEAPON[13],
        "Enabled": SCRAP_SWORD_WEAPON[14],
    }
    upsert_row(weapons, weapon_map, "Weapon ID", "weapon_scrap_sword", weapon_row)

    wb.save(WORKBOOK)
    print(f"Updated {WORKBOOK}")


if __name__ == "__main__":
    main()
