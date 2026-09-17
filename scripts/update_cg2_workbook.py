#!/usr/bin/env python3
"""Author CG2 appearance morph mapping sheet for human male/female variants."""

from __future__ import annotations

import openpyxl
from pathlib import Path

WORKBOOK = Path(__file__).resolve().parents[1] / "Chasma Design.xlsx"

MAPPING_HEADERS = [
    "Profile ID",
    "Variant ID",
    "Param ID",
    "Target Name",
    "Side",
    "Multiplier",
    "Enabled",
]

# Same semantic→technical mapping for both human variants (target names match GLB contract).
HUMAN_VARIANTS = ("human_male", "human_female")
HUMAN_MORPH_MAPPINGS = [
    ("build", "build_broad", "Above Default", 1.0),
    ("build", "build_narrow", "Below Default", 1.0),
    ("fat", "fat_soft", "Above Default", 1.0),
    ("muscle", "muscle_define", "Above Default", 1.0),
    ("head_size", "head_large", "Above Default", 1.0),
    ("head_size", "head_small", "Below Default", 1.0),
]


def ensure_sheet(wb, name: str, headers: list[str]):
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    return ws


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = ensure_sheet(wb, "Appearance Morph Mappings", MAPPING_HEADERS)
    rows: list[list[object]] = []
    for variant_id in HUMAN_VARIANTS:
        for param_id, target_name, side, multiplier in HUMAN_MORPH_MAPPINGS:
            rows.append(
                ["human", variant_id, param_id, target_name, side, multiplier, "Y"]
            )
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(WORKBOOK)
    print(f"Updated {WORKBOOK} with {len(rows)} CG2 morph mappings")


if __name__ == "__main__":
    main()
