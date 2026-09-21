#!/usr/bin/env python3
"""Author CG1 appearance profile sheets and unit linkage columns."""

from __future__ import annotations

import openpyxl
from pathlib import Path

WORKBOOK = Path(__file__).resolve().parents[1] / "Chasma Design.xlsx"

PROFILE_HEADERS = [
    "Profile ID",
    "Species ID",
    "Enabled",
    "Schema Version",
    "Height Min",
    "Height Max",
    "Height Default",
]

VARIANT_HEADERS = [
    "Profile ID",
    "Variant ID",
    "Render Key",
    "Display Name",
    "Enabled",
]

PARAMETER_HEADERS = [
    "Profile ID",
    "Param ID",
    "Display Name",
    "Category",
    "Min",
    "Max",
    "Default",
    "Display Order",
    "Enabled",
]

HUMAN_PARAMETERS = [
    ("build", "Build", "Body", 0.0, 1.0, 0.5, 1),
    ("fat", "Fat", "Body", 0.0, 1.0, 0.35, 2),
    ("muscle", "Muscle", "Body", 0.0, 1.0, 0.45, 3),
    ("head_size", "Head Size", "Head", 0.0, 1.0, 0.5, 4),
    ("shoulders", "Shoulders", "Regional", 0.0, 1.0, 0.5, 5),
    ("torso", "Torso", "Regional", 0.0, 1.0, 0.5, 6),
    ("arms", "Arms", "Regional", 0.0, 1.0, 0.5, 7),
    ("hips", "Hips", "Regional", 0.0, 1.0, 0.5, 8),
    ("legs", "Legs", "Regional", 0.0, 1.0, 0.5, 9),
]


def header_map(ws) -> dict[str, int]:
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    return {str(h).strip(): idx for idx, h in enumerate(headers) if h}


def ensure_column(ws, mapping: dict[str, int], name: str) -> int:
    if name in mapping:
        return mapping[name]
    col = ws.max_column
    ws.cell(row=1, column=col + 1, value=name)
    mapping[name] = col
    return col


def ensure_sheet(wb, name: str, headers: list[str]):
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
        for col, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=header)
    return ws


def write_rows(ws, headers: list[str], rows: list[list[object]]) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)

    profiles = ensure_sheet(wb, "Appearance Profiles", PROFILE_HEADERS)
    write_rows(
        profiles,
        PROFILE_HEADERS,
        [
            ["human", "human", "Y", 1, 0.85, 1.15, 1.0],
        ],
    )

    variants = ensure_sheet(wb, "Appearance Body Variants", VARIANT_HEADERS)
    write_rows(
        variants,
        VARIANT_HEADERS,
        [
            ["human", "human_male", "human_male", "Human Male", "Y"],
            ["human", "human_female", "human_female", "Human Female", "Y"],
        ],
    )

    parameters = ensure_sheet(wb, "Appearance Parameters", PARAMETER_HEADERS)
    write_rows(
        parameters,
        PARAMETER_HEADERS,
        [
            ["human", param_id, display, category, min_v, max_v, default, order, "Y"]
            for param_id, display, category, min_v, max_v, default, order in HUMAN_PARAMETERS
        ],
    )

    units = wb["Units"]
    umap = header_map(units)
    profile_col = ensure_column(units, umap, "Appearance Profile ID")
    variant_col = ensure_column(units, umap, "Default Body Variant ID")

    unit_links = {
        "U-0004": ("human", "human_male"),
        "U-0005": ("human", "human_female"),
    }
    id_col = umap["Unit ID"]
    for row in units.iter_rows(min_row=2):
        unit_id = row[id_col].value
        if unit_id not in unit_links:
            continue
        profile_id, variant_id = unit_links[unit_id]
        row[profile_col].value = profile_id
        row[variant_col].value = variant_id

    wb.save(WORKBOOK)
    print(f"Updated {WORKBOOK} for CG1 appearance foundation")


if __name__ == "__main__":
    main()
