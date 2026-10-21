#!/usr/bin/env py -3
"""Ensure bufomorph is importable: Species row + complete Units stats (U-0029)."""

from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "Chasma Design.xlsx"

PLACEHOLDER = {
    "Level": 3,
    "Base HP": 40,
    "Strength": 7,
    "Dexterity": 6,
    "Constitution": 6,
    "Agility": 5,
    "Charisma": 2,
    "Intelligence": 2,
    "Move Speed": 3.5,
    "Collision Radius": 1.0,
    "Max Slope": 45,
    "Sight Range": 20,
    "Rotation Correction Y Deg": 180,
    "Turn Speed Deg/s": 360,
}


def total_stats(row: dict) -> float:
    return sum(float(row[k]) for k in ("Strength", "Dexterity", "Constitution", "Agility", "Charisma", "Intelligence"))


def power_rating(level: float, base_hp: float, stats: float) -> float:
    return round(level * 2 + base_hp * 0.5 + stats * 0.8, 1)


def tier_label(rating: float) -> str:
    if rating >= 25:
        return "Elite"
    if rating >= 15:
        return "Veteran"
    if rating >= 8:
        return "Regular"
    return "Rookie"


def header_map(ws):
    return {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def upsert(ws, headers, key_col, key_val, data):
    kc = headers[key_col]
    row = None
    for r in range(2, ws.max_row + 2):
        v = ws.cell(r, kc).value
        if v is not None and str(v).strip() == key_val:
            row = r
            break
        if v is None or str(v).strip() == "":
            row = r
            break
    assert row is not None
    for col, val in data.items():
        if col in headers:
            ws.cell(row, headers[col], val)


def main():
    wb = openpyxl.load_workbook(WORKBOOK)
    species = wb["Species"]
    sh = header_map(species)
    upsert(species, sh, "Species Key", "bufomorph", {"Species Key": "bufomorph", "Name": "Bufomorph", "Enabled": "Y"})

    units = wb["Units"]
    uh = header_map(units)
    row_data = {
        "Unit ID": "U-0029",
        "Name": "Bufomorph",
        "Faction Key": "wild",
        "Species Key": "bufomorph",
        "File Path": r"assets\units\bufomorph.glb",
        "Animation Profile": "bufomorph",
        "Default Weapon ID": "weapon_bufomorph_primary",
        "Enabled": "Y",
        **PLACEHOLDER,
    }
    stats_total = total_stats(row_data)
    rating = power_rating(float(row_data["Level"]), float(row_data["Base HP"]), stats_total)
    if "Total Stats" in uh:
        row_data["Total Stats"] = stats_total
    if "Power Rating" in uh:
        row_data["Power Rating"] = rating
    if "Tier" in uh:
        row_data["Tier"] = tier_label(rating)
    upsert(units, uh, "Unit ID", "U-0029", row_data)
    wb.save(WORKBOOK)
    print(f"Updated {WORKBOOK}")


if __name__ == "__main__":
    main()
