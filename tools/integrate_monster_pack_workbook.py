#!/usr/bin/env py -3
"""Add monster-pack GLB units to Chasma Design.xlsx from import_manifest.json."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "Chasma Design.xlsx"
DEFAULT_MANIFEST = ROOT / "tools" / "monster_pack" / "import_manifest.json"
BACKUP = ROOT / "Chasma Design.xlsx.monster_pack.bak"

ANIMATION_OPTIONAL_COLUMNS = [
    "Death Animation",
    "Hit Reaction Animation",
    "Upper Body Split Bone",
    "Turn Left Animation",
    "Turn Right Animation",
    "Turn Left Duration",
    "Turn Right Duration",
]

# PLACEHOLDER gameplay stats — no design source for pack monsters.
PLACEHOLDER_UNIT_STATS = {
    "Level": 3,
    "Base HP": 40,
    "Strength": 7,
    "Dexterity": 6,
    "Constitution": 6,
    "Agility": 5,
    "Charisma": 2,
    "Intelligence": 2,
    "Move Speed": 3.5,
    "Collision Radius": 1.0,
    "Max Slope": 45,
    "Sight Range": 20,
    "Rotation Correction Y Deg": 180,
    "Turn Speed Deg/s": 360,
}

ATTACK_CLIP_PRIORITY = [
    re.compile(r"attack", re.I),
    re.compile(r"bite", re.I),
    re.compile(r"claw", re.I),
    re.compile(r"smash", re.I),
    re.compile(r"combo", re.I),
    re.compile(r"spit", re.I),
    re.compile(r"sting", re.I),
    re.compile(r"slash", re.I),
]


def header_map(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }


def ensure_columns(ws, columns: list[str]) -> dict[str, int]:
    headers = header_map(ws)
    next_col = ws.max_column + 1
    for column in columns:
        if column not in headers:
            ws.cell(1, next_col, column)
            headers[column] = next_col
            next_col += 1
    return headers


def total_stats(row: dict[str, object]) -> float:
    return sum(float(row[name]) for name in ("Strength", "Dexterity", "Constitution", "Agility", "Charisma", "Intelligence"))


def power_rating(level: float, base_hp: float, stats: float) -> float:
    return round(level * 2 + base_hp * 0.5 + stats * 0.8, 1)


def tier_label(rating: float) -> str:
    if rating >= 25:
        return "Elite"
    if rating >= 15:
        return "Veteran"
    if rating >= 8:
        return "Regular"
    return "Rookie"


def upsert_row_by_key(ws, headers: dict[str, int], key_column: str, key_value: str, data: dict) -> None:
    key_col = headers[key_column]
    target_row = None
    for row in range(2, ws.max_row + 2):
        existing = ws.cell(row, key_col).value
        if existing is not None and str(existing).strip() == key_value:
            target_row = row
            break
        if existing is None or str(existing).strip() == "":
            target_row = row
            break
    assert target_row is not None
    for column, value in data.items():
        if column not in headers:
            continue
        ws.cell(target_row, headers[column], value)


def next_unit_id(ws, headers: dict[str, int]) -> str:
    col = headers["Unit ID"]
    max_n = 0
    for row in range(2, ws.max_row + 1):
        raw = ws.cell(row, col).value
        if not raw:
            continue
        m = re.match(r"U-(\d+)$", str(raw).strip())
        if m:
            max_n = max(max_n, int(m.group(1)))
    return f"U-{max_n + 1:04d}"


def pick_clip(clips: list[str], candidates: list[str]) -> str:
    clip_set = set(clips)
    for name in candidates:
        if name in clip_set:
            return name
    return ""


def pick_non_rm(clips: list[str], candidates: list[str]) -> str:
    for name in candidates:
        if name in clips:
            return name
        rm = f"{name}_RM"
        if rm in clips and name not in clips:
            continue
    return pick_clip(clips, candidates)


def build_animation_profile(render_key: str, clips: list[str]) -> dict:
    idle = pick_clip(
        clips,
        ["Idle", "idle", "IdleBreathe", "Idle1", "Idle_01", "IdleBreath"],
    )
    walk = pick_non_rm(
        clips,
        [
            "CrawlForward",
            "Walk",
            "WalkForward",
            "Move",
            "RunForward",
            "Locomotion",
        ],
    )
    run = pick_non_rm(clips, ["Run", "CrawlForward", "Walk", "WalkForward"])
    death = pick_clip(clips, ["Death", "Die"])
    hit = pick_clip(clips, ["GetHit1", "GetHit", "Hit", "TakeHit"])
    turn_l = pick_non_rm(clips, ["CrawlLeft", "TurnLeft", "LeftTurn"])
    turn_r = pick_non_rm(clips, ["CrawlRight", "TurnRight", "RightTurn"])
    ref_speed = PLACEHOLDER_UNIT_STATS["Move Speed"]
    return {
        "Profile ID": render_key,
        "Idle Animation": idle or "Idle",
        "Walk Animation": walk,
        "Run Animation": run,
        "Locomotion Reference Speed": ref_speed,
        "Enabled": "Y",
        "Death Animation": death,
        "Hit Reaction Animation": hit,
        "Upper Body Split Bone": "",
        "Turn Left Animation": turn_l,
        "Turn Right Animation": turn_r,
        "Turn Left Duration": 1.0 if turn_l else "",
        "Turn Right Duration": 1.0 if turn_r else "",
    }


def pick_attack_clip(clips: list[str]) -> str | None:
    usable = [c for c in clips if not c.endswith("_RM") and "|" not in c and "Take 001" not in c]
    for pattern in ATTACK_CLIP_PRIORITY:
        for clip in usable:
            if pattern.search(clip):
                return clip
    return None


def build_weapon(render_key: str, display_name: str, attack_clip: str) -> dict:
    weapon_id = f"weapon_{render_key}_primary"
    return {
        "Weapon ID": weapon_id,
        "Name": f"{display_name} Attack",
        "Description": "Intrinsic monster attack (placeholder balance).",
        "Damage": 6,
        "Damage Type": "Slashing",
        "Range": 1.5,
        "Attacks Per Second": 1.0,
        "Windup": 0.2,
        "Recovery": 0.15,
        "Hit Mode": "Melee",
        "Projectile Key": None,
        "Animation Key": attack_clip,
        "Target Filters": "Enemies, Wildlife",
        "Stat Scaling": None,
        "Enabled": "Y",
    }


def build_unit_row(
    unit_id: str,
    display_name: str,
    render_key: str,
    weapon_id: str | None,
    headers: dict[str, int],
) -> dict:
    row = {
        "Unit ID": unit_id,
        "Name": display_name,
        "Faction Key": "wild",
        "Species Key": render_key,
        "File Path": rf"assets\units\{render_key}.glb",
        "Animation Profile": render_key,
        "Default Weapon ID": weapon_id or "",
        "Enabled": "Y",
    }
    row.update(PLACEHOLDER_UNIT_STATS)
    stats_total = total_stats(row)
    rating = power_rating(float(row["Level"]), float(row["Base HP"]), stats_total)
    if "Total Stats" in headers:
        row["Total Stats"] = stats_total
    if "Power Rating" in headers:
        row["Power Rating"] = rating
    if "Tier" in headers:
        row["Tier"] = tier_label(rating)
    return row


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if not WORKBOOK.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK}")
    if not args.manifest.is_file():
        raise SystemExit(f"Manifest not found: {args.manifest} (run batch export first)")

    entries = json.loads(args.manifest.read_text(encoding="utf-8-sig"))
    if not entries:
        raise SystemExit("Manifest is empty")

    if args.dry_run:
        print(f"Would integrate {len(entries)} monsters from {args.manifest}")
        return

    shutil.copy2(WORKBOOK, BACKUP)
    wb = openpyxl.load_workbook(WORKBOOK)

    species_ws = wb["Species"]
    species_headers = header_map(species_ws)
    anim_ws = wb["Animation Profiles"]
    anim_headers = ensure_columns(anim_ws, ANIMATION_OPTIONAL_COLUMNS)
    weapons_ws = wb["Weapons"]
    weapon_headers = header_map(weapons_ws)
    units_ws = wb["Units"]
    unit_headers = ensure_columns(
        units_ws,
        [
            "Default Weapon ID",
            "Rotation Correction X Deg",
            "Rotation Correction Y Deg",
            "Rotation Correction Z Deg",
            "Turn Speed Deg/s",
        ],
    )

    next_id = next_unit_id(units_ws, unit_headers)

    for entry in entries:
        render_key = entry["render_key"]
        display_name = entry.get("monster") or render_key.replace("_", " ").title()
        clips = list(entry.get("clips") or [])

        upsert_row_by_key(
            species_ws,
            species_headers,
            "Species Key",
            render_key,
            {"Species Key": render_key, "Name": display_name, "Enabled": "Y"},
        )

        profile = build_animation_profile(render_key, clips)
        upsert_row_by_key(anim_ws, anim_headers, "Profile ID", render_key, profile)

        attack = pick_attack_clip(clips)
        if attack:
            weapon = build_weapon(render_key, display_name, attack)
            weapon_id = weapon["Weapon ID"]
            upsert_row_by_key(weapons_ws, weapon_headers, "Weapon ID", weapon_id, weapon)
        else:
            # Workbook requires Default Weapon ID when column exists; no attack clip mapped.
            weapon_id = "weapon_fists"

        unit_row = build_unit_row(next_id, display_name, render_key, weapon_id, unit_headers)
        upsert_row_by_key(units_ws, unit_headers, "Unit ID", next_id, unit_row)
        print(f"{next_id} {display_name} ({render_key}) weapon={weapon_id or 'none'}")
        m = re.match(r"U-(\d+)$", next_id)
        next_id = f"U-{int(m.group(1)) + 1:04d}" if m else next_id

    wb.save(WORKBOOK)
    print(f"Updated {WORKBOOK}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
